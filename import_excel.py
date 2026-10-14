import os
import sqlite3
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "scada.db")
EXCEL_FILE = os.path.join(BASE_DIR, "scada_data.xlsx")  # మీ ఎక్సెల్ ఫైల్ పేరు

def import_excel_to_db():
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Error: '{EXCEL_FILE}' ఫైల్ దొరకలేదు! ఫైల్ పేరు సరిచూసుకోండి.")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 1. టేబుల్స్ రీసెట్ & క్రియేషన్
    cursor.execute("DROP TABLE IF EXISTS assets")
    cursor.execute("DROP TABLE IF EXISTS pipelines")

    cursor.execute("""
        CREATE TABLE assets (
            id TEXT PRIMARY KEY,
            name TEXT,
            type TEXT,
            tag TEXT,
            details TEXT,
            status TEXT DEFAULT 'ON',
            location TEXT,
            pressure REAL DEFAULT 0.0
        )
    """)

    cursor.execute("""
        CREATE TABLE pipelines (
            source_id TEXT,
            target_id TEXT,
            PRIMARY KEY (source_id, target_id)
        )
    """)

    # 2. Assets Sheet ఇంపోర్ట్
    if "Assets" in wb.sheetnames:
        sheet = wb["Assets"]
        headers = [cell.value for cell in sheet[1]]
        
        # హెడర్ ఇండెక్స్‌లు గుర్తించడం
        col_map = {str(name).strip(): idx for idx, name in enumerate(headers) if name}
        
        assets_data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            asset_id = str(row[col_map.get("Asset ID", 0)] or "").strip()
            if not asset_id:
                continue

            name = str(row[col_map.get("Asset Name", 1)] or asset_id).strip()
            a_type = str(row[col_map.get("Type", 2)] or "VALVE").strip().upper()
            tag = str(row[col_map.get("Tag", 3)] or "125 MM").strip()
            details = str(row[col_map.get("Details", 4)] or "").strip()
            status = str(row[col_map.get("Status", 7)] or "ON").strip().upper()
            location = str(row[col_map.get("Location", 8)] or "").strip()
            
            # DRS అయితే Details నుండి ప్రెషర్ లేదా డీఫాల్ట్ ప్రెషర్ కేటాయించడం
            pressure = 4.0 if a_type == "DRS" else 0.0
            if a_type == "DRS" and "bar" in details.lower():
                try:
                    pressure = float(''.join(c for c in details if c.isdigit() or c == '.'))
                except:
                    pass

            assets_data.append((asset_id, name, a_type, tag, details, status, location, pressure))

        cursor.executemany("INSERT INTO assets VALUES (?,?,?,?,?,?,?,?)", assets_data)
        print(f"✅ {len(assets_data)} Assets విజయవంతంగా ఇంపోర్ట్ అయ్యాయి.")
    else:
        print("⚠️ 'Assets' షీట్ కనపడలేదు!")

    # 3. Links Sheet ఇంపోర్ట్
    links_sheet_name = "Links" if "Links" in wb.sheetnames else "Sheet2"
    if links_sheet_name in wb.sheetnames:
        sheet = wb[links_sheet_name]
        headers = [cell.value for cell in sheet[1]]
        col_map = {str(name).strip(): idx for idx, name in enumerate(headers) if name}

        parent_idx = col_map.get("Parent ID", 0)
        child_idx = col_map.get("Child ID", 1)

        links_data = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            src = str(row[parent_idx] or "").strip()
            tgt = str(row[child_idx] or "").strip()
            if src and tgt and src != tgt:
                links_data.add((src, tgt))

        cursor.executemany("INSERT OR IGNORE INTO pipelines VALUES (?,?)", list(links_data))
        print(f"✅ {len(links_data)} Pipeline Links విజయవంతంగా ఇంపోర్ట్ అయ్యాయి.")
    else:
        print("⚠️ 'Links' షీట్ కనపడలేదు!")

    conn.commit()
    conn.close()
    print("🚀 'scada.db' డేటాబేస్ పర్ఫెక్ట్‌గా అప్‌డేట్ అయింది!")

if __name__ == "__main__":
    import_excel_to_db()