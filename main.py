import os
import io
import json
import zipfile
import sqlite3
import xml.etree.ElementTree as ET
from datetime import datetime
import openpyxl
from typing import Dict, List, Any
from fastapi import FastAPI, UploadFile, File
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
import networkx as nx

app = FastAPI(title="BGL Gas SCADA Engine v7.0")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INDEX_HTML = os.path.join(BASE_DIR, "templates", "index.html")
DB_PATH = os.path.join(BASE_DIR, "scada.db")
MASTER_EXCEL = os.path.join(BASE_DIR, "master_data.xlsx")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def load_excel_bytes_to_db(contents: bytes, cursor):
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    now = datetime.now().strftime("%Y-%m-%d")

    cursor.execute("DELETE FROM assets")
    cursor.execute("DELETE FROM pipelines")

    assets_sheet = None
    for s in wb.sheetnames:
        if s.strip().lower() == "assets":
            assets_sheet = wb[s]
            break
    if not assets_sheet:
        assets_sheet = wb.active

    asset_count = 0
    if assets_sheet:
        headers = [str(c.value or "").strip() for c in assets_sheet[1]]
        col_map = {name: idx for idx, name in enumerate(headers) if name}
        rows = []
        for row in assets_sheet.iter_rows(min_row=2, values_only=True):
            if not row or len(row) == 0:
                continue
            raw_id = row[col_map.get("Asset ID", 0)]
            if raw_id is None or str(raw_id).strip() == "":
                continue
            aid = str(raw_id).strip()
            name = str(row[col_map.get("Asset Name", 1)] or aid).strip()
            atype = str(row[col_map.get("Type", 2)] or "VALVE").strip().upper()
            tag = str(row[col_map.get("Tag", 3)] or "125 MM").strip()
            details = str(row[col_map.get("Details", 4)] or "").strip() if len(row) > 4 else ""
            status = str(row[col_map.get("Status", 7)] or "ON").strip().upper() if len(row) > 7 else "ON"
            if status not in ["ON", "OFF"]:
                status = "ON"
            loc = str(row[col_map.get("Location", 8)] or "").strip() if len(row) > 8 else ""
            if loc == "None":
                loc = ""

            press = 4.0 if atype == "DRS" else 0.0
            if atype == "DRS" and "bar" in details.lower():
                try:
                    press = float("".join(ch for ch in details if ch.isdigit() or ch == "."))
                except:
                    pass
            rows.append((aid, name, atype, tag, details, now, now, status, loc, press))
        cursor.executemany("INSERT OR REPLACE INTO assets VALUES (?,?,?,?,?,?,?,?,?,?)", rows)
        asset_count = len(rows)

    links_sheet = None
    for s in wb.sheetnames:
        if s.strip().lower() in ["links", "link", "sheet2"]:
            links_sheet = wb[s]
            break

    link_count = 0
    if links_sheet:
        headers = [str(c.value or "").strip() for c in links_sheet[1]]
        col_map = {name: idx for idx, name in enumerate(headers) if name}
        p_idx = col_map.get("Parent ID", 0)
        c_idx = col_map.get("Child ID", 1)
        links = []
        for row in links_sheet.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 2:
                continue
            src = str(row[p_idx] or "").strip() if row[p_idx] is not None else ""
            tgt = str(row[c_idx] or "").strip() if row[c_idx] is not None else ""
            if src and tgt and src != tgt and src != "None" and tgt != "None":
                links.append((src, tgt, now))
        cursor.executemany("INSERT OR IGNORE INTO pipelines VALUES (?,?,?)", links)
        link_count = len(links)

    return asset_count, link_count

def seed_default_network(cursor):
    if os.path.exists(MASTER_EXCEL):
        try:
            with open(MASTER_EXCEL, "rb") as f:
                load_excel_bytes_to_db(f.read(), cursor)
            return
        except Exception as e:
            print("Master excel load error:", e)

    now = datetime.now().strftime("%Y-%m-%d")
    default_assets = [
        ("D01", "MS DRS", "DRS", "125 MM", "4.1 Bar", now, now, "ON", "17.6231,78.5253", 4.1),
        ("D03", "Check Post DRS", "DRS", "125 MM", "3.6 Bar", now, now, "ON", "17.6105,78.4966", 3.6),
        ("D02", "Depo DRS", "DRS", "125 MM", "3.1 Bar", now, now, "ON", "17.6313,78.5299", 3.1),
        ("CG-D01-PE01", "CG-D01-PE01", "VALVE", "125 MM", "", now, now, "ON", "17.6231,78.5253", 0.0),
        ("CG-D01-PE02", "CG-D01-PE02", "VALVE", "125 MM", "", now, now, "ON", "17.6343,78.5177", 0.0),
        ("31000003", "IPE Univ. Canteen", "CPNG", "125 MM", "", now, now, "ON", "17.6240,78.5260", 0.0),
        ("S03-D03-PE01", "S03-D03-PE01", "VALVE", "125 MM", "", now, now, "ON", "17.6105,78.4966", 0.0),
        ("S02-D02-PE01", "S02-D02-PE01", "VALVE", "125 MM", "", now, now, "ON", "17.6313,78.5299", 0.0)
    ]
    cursor.executemany("INSERT OR REPLACE INTO assets VALUES (?,?,?,?,?,?,?,?,?,?)", default_assets)

    default_pipelines = [
        ("D01", "CG-D01-PE01", now), ("CG-D01-PE01", "31000003", now),
        ("D01", "CG-D01-PE02", now), ("D03", "S03-D03-PE01", now),
        ("D02", "S02-D02-PE01", now), ("CG-D01-PE02", "S03-D03-PE01", now)
    ]
    cursor.executemany("INSERT OR REPLACE INTO pipelines VALUES (?,?,?)", default_pipelines)

def init_db():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS assets (
            id TEXT PRIMARY KEY,
            name TEXT,
            type TEXT,
            tag TEXT,
            details TEXT,
            created_at TEXT,
            updated_at TEXT,
            status TEXT DEFAULT 'ON',
            location TEXT,
            pressure REAL DEFAULT 0.0
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS pipelines (
            source_id TEXT,
            target_id TEXT,
            created_at TEXT,
            PRIMARY KEY (source_id, target_id)
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            username TEXT PRIMARY KEY,
            password TEXT,
            role TEXT DEFAULT 'user'
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS kmz_overlays (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT DEFAULT 'KMZ_File',
            name TEXT,
            geom_type TEXT,
            coords_json TEXT
        )
    """)

    # Ensure file_name column exists in kmz_overlays
    cursor.execute("PRAGMA table_info(kmz_overlays)")
    kcols = [r[1] for r in cursor.fetchall()]
    if "file_name" not in kcols:
        cursor.execute("DROP TABLE IF EXISTS kmz_overlays")
        cursor.execute("""
            CREATE TABLE kmz_overlays (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                file_name TEXT DEFAULT 'KMZ_File',
                name TEXT,
                geom_type TEXT,
                coords_json TEXT
            )
        """)
        conn.commit()

    cursor.execute("SELECT COUNT(*) FROM users")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO users VALUES ('admin', 'admin123', 'admin')")
        cursor.execute("INSERT INTO users VALUES ('user', 'user123', 'user')")
        conn.commit()

    cursor.execute("PRAGMA table_info(assets)")
    cols = [r[1] for r in cursor.fetchall()]
    if "created_at" not in cols or "pressure" not in cols:
        cursor.execute("DROP TABLE IF EXISTS assets")
        cursor.execute("DROP TABLE IF EXISTS pipelines")
        conn.commit()
        conn.close()
        init_db()
        return

    cursor.execute("SELECT COUNT(*) FROM assets")
    if cursor.fetchone()[0] == 0:
        seed_default_network(cursor)
        conn.commit()
    conn.close()

init_db()

class LoginReq(BaseModel):
    username: str
    password: str

class UserSaveReq(BaseModel):
    username: str
    password: str
    role: str = "user"

class UserDeleteReq(BaseModel):
    username: str

class ToggleReq(BaseModel):
    asset_id: str
    status: str

class SimulateReq(BaseModel):
    overrides: Dict[str, str]

class PressureReq(BaseModel):
    drs_id: str
    pressure: float

class AssetSaveReq(BaseModel):
    id: str
    name: str
    type: str
    tag: str = "125 MM"
    details: str = ""
    status: str = "ON"
    location: str = ""
    pressure: float = 0.0
    parent_id: str = ""

class AssetDeleteReq(BaseModel):
    id: str

class PipelineReq(BaseModel):
    source_id: str
    target_id: str

class KmzDeleteReq(BaseModel):
    id: int = 0
    file_name: str = ""

class BackupRestoreReq(BaseModel):
    assets: Dict[str, Any]
    pipelines: List[List[str]]
    kmz_overlays: List[Dict[str, Any]] = []

def is_upstream_node(nid: str, assets: dict) -> bool:
    u = nid.strip().upper()
    if u == "CG" or (u.startswith("S0") and len(u) == 3):
        return True
    t = str(assets.get(nid, {}).get("type", "")).upper()
    if t in ["CG", "SV", "STEEL"]:
        return True
    return False

def load_network_from_db(overrides: Dict[str, str] = None):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM assets")
    assets = {row["id"]: dict(row) for row in cursor.fetchall()}
    cursor.execute("SELECT source_id, target_id FROM pipelines")
    pipelines = [(row["source_id"], row["target_id"]) for row in cursor.fetchall()]
    cursor.execute("SELECT id, file_name, name, geom_type, coords_json FROM kmz_overlays ORDER BY id DESC")
    kmz_list = [dict(row) for row in cursor.fetchall()]
    conn.close()

    if overrides:
        for aid, st in overrides.items():
            if aid in assets:
                assets[aid]["status"] = st

    G = nx.Graph()
    for node_id in assets:
        if not is_upstream_node(node_id, assets):
            G.add_node(node_id)
    for u, v in pipelines:
        if u in G and v in G and not is_upstream_node(u, assets) and not is_upstream_node(v, assets):
            G.add_edge(u, v)

    return assets, pipelines, kmz_list, G

def run_scada_bfs_flow(overrides: Dict[str, str] = None):
    assets, pipelines, kmz_list, G = load_network_from_db(overrides)

    active_drs = [
        k for k, v in assets.items()
        if str(v.get("type", "")).upper() == "DRS" and str(v.get("status", "ON")).upper() == "ON"
    ]
    active_drs.sort(key=lambda x: float(assets[x].get("pressure") or 0.0), reverse=True)

    gas_supply = {node: False for node in assets}
    fed_by = {}
    flow_rank = {}
    step_counter = 1

    for drs in active_drs:
        press = float(assets[drs].get("pressure") or 4.0)
        gas_supply[drs] = True
        fed_by[drs] = {"id": drs, "pressure": press}
        if drs not in flow_rank:
            flow_rank[drs] = step_counter
            step_counter += 1

        q = [drs]
        visited = {drs}

        while q:
            curr = q.pop(0)
            c_type = str(assets.get(curr, {}).get("type", "")).upper()
            c_stat = str(assets.get(curr, {}).get("status", "ON")).upper()

            if ("VALVE" in c_type or c_type == "VC") and c_stat == "OFF":
                gas_supply[curr] = False
                continue

            if curr in G:
                for nbr in G.neighbors(curr):
                    if nbr not in visited:
                        n_type = str(assets.get(nbr, {}).get("type", "")).upper()
                        n_stat = str(assets.get(nbr, {}).get("status", "ON")).upper()
                        if ("VALVE" in n_type or n_type == "VC") and n_stat == "OFF":
                            visited.add(nbr)
                            gas_supply[nbr] = False
                            continue
                        visited.add(nbr)
                        gas_supply[nbr] = True
                        if nbr not in fed_by:
                            fed_by[nbr] = fed_by[curr]
                        if nbr not in flow_rank:
                            flow_rank[nbr] = step_counter
                            step_counter += 1
                        q.append(nbr)

    return assets, pipelines, kmz_list, gas_supply, fed_by, flow_rank

def build_response(msg="", overrides: Dict[str, str] = None):
    assets, pipelines, kmz_list, supply, fed_by, flow_rank = run_scada_bfs_flow(overrides)
    return {
        "success": True,
        "message": msg,
        "assets": assets,
        "pipelines": pipelines,
        "kmz_overlays": kmz_list,
        "supply": supply,
        "fed_by": fed_by,
        "flow_rank": flow_rank
    }

@app.get("/")
def index():
    if os.path.exists(INDEX_HTML):
        return FileResponse(INDEX_HTML, media_type="text/html")
    return {"error": "index.html not found"}

@app.post("/api/login")
def api_login(req: LoginReq):
    u = req.username.strip()
    p = req.password.strip()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT username, role FROM users WHERE LOWER(username) = LOWER(?) AND password = ?", (u, p))
    row = cursor.fetchone()
    conn.close()
    if row:
        return {"success": True, "username": row["username"], "role": row["role"]}
    return {"success": False, "message": "Invalid Login ID leda Password!"}

@app.get("/api/users")
def list_users():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT username, role FROM users ORDER BY role, username")
    users = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return {"success": True, "users": users}

@app.post("/api/users/save")
def save_user(req: UserSaveReq):
    u = req.username.strip()
    p = req.password.strip()
    r = req.role.strip().lower()
    if not u or not p:
        return {"success": False, "message": "Username mariyu Password tappanisari!"}
    if r not in ["admin", "user"]:
        r = "user"
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO users (username, password, role) VALUES (?, ?, ?)", (u, p, r))
    conn.commit()
    cursor.execute("SELECT username, role FROM users ORDER BY role, username")
    users = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"success": True, "message": f"Login ID '{u}' ({r.upper()}) saved!", "users": users}

@app.post("/api/users/delete")
def delete_user(req: UserDeleteReq):
    u = req.username.strip()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM users WHERE role = 'admin'")
    admin_count = cursor.fetchone()[0]
    cursor.execute("SELECT role FROM users WHERE username = ?", (u,))
    target = cursor.fetchone()
    if target and target["role"] == "admin" and admin_count <= 1:
        conn.close()
        return {"success": False, "message": "Kaneesam okka Admin ID aina undali!"}
    cursor.execute("DELETE FROM users WHERE username = ?", (u,))
    conn.commit()
    cursor.execute("SELECT username, role FROM users ORDER BY role, username")
    users = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"success": True, "message": f"Login ID '{u}' deleted!", "users": users}

@app.get("/api/snapshot")
def get_snapshot():
    return build_response()

@app.post("/api/simulate")
def simulate_flow(req: SimulateReq):
    return build_response("Simulation updated", req.overrides)

@app.post("/api/restore-backup")
def restore_backup(req: BackupRestoreReq):
    if not req.assets:
        return build_response()
    now = datetime.now().strftime("%Y-%m-%d")
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM assets")
    cursor.execute("DELETE FROM pipelines")

    a_rows = []
    for aid, a in req.assets.items():
        a_rows.append((
            str(aid),
            str(a.get("name") or aid),
            str(a.get("type") or "VALVE").upper(),
            str(a.get("tag") or "125 MM"),
            str(a.get("details") or ""),
            str(a.get("created_at") or now),
            str(a.get("updated_at") or now),
            str(a.get("status") or "ON").upper(),
            str(a.get("location") or ""),
            float(a.get("pressure") or 0.0)
        ))
    cursor.executemany("INSERT OR REPLACE INTO assets VALUES (?,?,?,?,?,?,?,?,?,?)", a_rows)

    p_rows = []
    for pair in req.pipelines:
        if len(pair) >= 2 and pair[0] and pair[1]:
            p_rows.append((str(pair[0]), str(pair[1]), now))
    cursor.executemany("INSERT OR IGNORE INTO pipelines VALUES (?,?,?)", p_rows)

    if req.kmz_overlays:
        cursor.execute("DELETE FROM kmz_overlays")
        k_rows = [(k.get("file_name", "KMZ_File"), k.get("name", ""), k.get("geom_type", "line"), k.get("coords_json", "[]")) for k in req.kmz_overlays]
        cursor.executemany("INSERT INTO kmz_overlays (file_name, name, geom_type, coords_json) VALUES (?,?,?,?)", k_rows)

    conn.commit()
    conn.close()
    return build_response("Backup data auto-restored!")

@app.post("/api/asset/save")
def save_asset(req: AssetSaveReq):
    asset_id = req.id.strip()
    if not asset_id:
        return {"success": False, "message": "Asset ID ivvali!"}

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    a_type = req.type.strip().upper()
    press = req.pressure if a_type == "DRS" else 0.0
    if a_type == "DRS" and press == 0.0:
        press = 4.0

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT created_at FROM assets WHERE id = ?", (asset_id,))
    existing = cursor.fetchone()
    created_at = existing["created_at"] if existing else now

    cursor.execute("""
        INSERT OR REPLACE INTO assets 
        (id, name, type, tag, details, created_at, updated_at, status, location, pressure)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        asset_id, req.name.strip() or asset_id, a_type,
        req.tag.strip() or "125 MM", req.details.strip(),
        created_at, now, req.status.strip().upper() or "ON",
        req.location.strip(), press
    ))

    if req.parent_id and req.parent_id.strip() and req.parent_id.strip() != asset_id:
        cursor.execute("INSERT OR IGNORE INTO pipelines VALUES (?, ?, ?)", (req.parent_id.strip(), asset_id, now))

    conn.commit()
    conn.close()
    return build_response(f"Asset '{asset_id}' saved permanently!")

@app.post("/api/asset/delete")
def delete_asset(req: AssetDeleteReq):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM assets WHERE id = ?", (req.id.strip(),))
    cursor.execute("DELETE FROM pipelines WHERE source_id = ? OR target_id = ?", (req.id.strip(), req.id.strip()))
    conn.commit()
    conn.close()
    return build_response("Asset deleted permanently!")

@app.post("/api/pipeline/add")
def add_pipeline(req: PipelineReq):
    src, tgt = req.source_id.strip(), req.target_id.strip()
    if not src or not tgt or src == tgt:
        return {"success": False, "message": "Valid Parent mariyu Child Asset ID enter cheyandi!"}
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("INSERT OR IGNORE INTO pipelines VALUES (?, ?, ?)", (src, tgt, now))
    conn.commit()
    conn.close()
    return build_response(f"Link '{src}' ➔ '{tgt}' added permanently!")

@app.post("/api/pipeline/delete")
def delete_pipeline(req: PipelineReq):
    src, tgt = req.source_id.strip(), req.target_id.strip()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM pipelines WHERE (source_id=? AND target_id=?) OR (source_id=? AND target_id=?)", (src, tgt, tgt, src))
    conn.commit()
    conn.close()
    return build_response(f"Link '{src}' ⟷ '{tgt}' removed permanently!")

@app.post("/api/toggle")
def toggle_asset(req: ToggleReq):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE assets SET status = ? WHERE id = ?", (req.status, req.asset_id))
    conn.commit()
    conn.close()
    return build_response()

@app.post("/api/pressure")
def update_pressure(req: PressureReq):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE assets SET pressure = ?, details = ? WHERE id = ?", (req.pressure, f"{req.pressure} Bar", req.drs_id))
    conn.commit()
    conn.close()
    return build_response()

@app.post("/api/reset")
def reset_all():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE assets SET status = 'ON'")
    conn.commit()
    conn.close()
    return build_response("All valves reset to ON permanently!")

@app.post("/api/upload-excel")
async def upload_excel(file: UploadFile = File(...)):
    try:
        contents = await file.read()
        with open(MASTER_EXCEL, "wb") as f:
            f.write(contents)
        conn = get_db()
        cursor = conn.cursor()
        asset_count, link_count = load_excel_bytes_to_db(contents, cursor)
        conn.commit()
        conn.close()
        return build_response(f"Excel Imported & Saved Permanently: {asset_count} Assets, {link_count} Links!")
    except Exception as e:
        return {"success": False, "message": f"Excel Read Error: {str(e)}"}

def parse_kml_coordinates(coord_text: str):
    pts = []
    for token in (coord_text or "").strip().split():
        parts = token.split(",")
        if len(parts) >= 2:
            try:
                lng = float(parts[0])
                lat = float(parts[1])
                pts.append([lat, lng])
            except:
                pass
    return pts

@app.post("/api/upload-kmz")
async def upload_kmz(file: UploadFile = File(...)):
    try:
        raw = await file.read()
        orig_filename = (file.filename or "KMZ_File").strip()
        fname = orig_filename.lower()
        kml_bytes = None
        if fname.endswith(".kmz") or zipfile.is_zipfile(io.BytesIO(raw)):
            with zipfile.ZipFile(io.BytesIO(raw), "r") as z:
                for name in z.namelist():
                    if name.lower().endswith(".kml"):
                        kml_bytes = z.read(name)
                        break
        else:
            kml_bytes = raw

        if not kml_bytes:
            return {"success": False, "message": "KMZ/KML file lo valid KML data kanipinchaledhu!"}

        root = ET.fromstring(kml_bytes)
        overlays = []

        for pm in root.iter():
            if pm.tag.endswith("Placemark"):
                pm_name = "Pipeline"
                for child in pm:
                    if child.tag.endswith("name") and child.text:
                        pm_name = child.text.strip()
                        break

                for elem in pm.iter():
                    if elem.tag.endswith("LineString") or elem.tag.endswith("LinearRing"):
                        for sub in elem.iter():
                            if sub.tag.endswith("coordinates") and sub.text:
                                pts = parse_kml_coordinates(sub.text)
                                if len(pts) >= 2:
                                    overlays.append((orig_filename, pm_name, "line", json.dumps(pts)))
                    elif elem.tag.endswith("Point"):
                        for sub in elem.iter():
                            if sub.tag.endswith("coordinates") and sub.text:
                                pts = parse_kml_coordinates(sub.text)
                                if len(pts) == 1:
                                    overlays.append((orig_filename, pm_name, "point", json.dumps(pts[0])))

        conn = get_db()
        cursor = conn.cursor()
        # Remove old layers from same filename if re-uploaded, keep other valid files
        cursor.execute("DELETE FROM kmz_overlays WHERE file_name = ?", (orig_filename,))
        cursor.executemany("INSERT INTO kmz_overlays (file_name, name, geom_type, coords_json) VALUES (?,?,?,?)", overlays)
        conn.commit()
        conn.close()
        return build_response(f"'{orig_filename}' Imported ({len(overlays)} Map Lines/Points added)!")
    except Exception as e:
        return {"success": False, "message": f"KMZ Parse Error: {str(e)}"}

@app.post("/api/kmz/delete")
def delete_kmz_item(req: KmzDeleteReq):
    conn = get_db()
    cursor = conn.cursor()
    if req.file_name:
        cursor.execute("DELETE FROM kmz_overlays WHERE file_name = ?", (req.file_name,))
        msg = f"KMZ File '{req.file_name}' deleted!"
    else:
        cursor.execute("DELETE FROM kmz_overlays WHERE id = ?", (req.id,))
        msg = "KMZ Layer deleted!"
    conn.commit()
    conn.close()
    return build_response(msg)

@app.post("/api/kmz/clear")
def clear_all_kmz():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM kmz_overlays")
    conn.commit()
    conn.close()
    return build_response("All KMZ files & map lines deleted!")

@app.get("/api/export-excel")
def export_excel():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM assets")
    ar = cursor.fetchall()
    cursor.execute("SELECT * FROM pipelines")
    lr = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws_a = wb.active
    ws_a.title = "Assets"
    ws_a.append(["Asset ID", "Asset Name", "Type", "Tag", "Details", "Created At", "Updated At", "Status", "Location"])
    for r in ar:
        ws_a.append([r["id"], r["name"], r["type"], r["tag"], r["details"], r["created_at"], r["updated_at"], r["status"], r["location"]])

    ws_l = wb.create_sheet(title="Links")
    ws_l.append(["Parent ID", "Child ID", "Created At"])
    for r in lr:
        ws_l.append([r["source_id"], r["target_id"], r["created_at"]])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=BGL_SCADA_Data.xlsx"})